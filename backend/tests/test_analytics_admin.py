import io

from openpyxl import load_workbook


def test_analytics(client, user_h, viewer_h):
    for v in (10, 20, 30):
        client.post("/api/v1/records", json={"title": "s", "value": v, "category": "stats", "timestamp": "2026-03-01T10:00:00"}, headers=user_h)
    client.post("/api/v1/records", json={"title": "s", "value": 40, "category": "stats", "timestamp": "2026-03-02T10:00:00"}, headers=user_h)

    s = client.get("/api/v1/analytics/summary", params={"category": "stats"}, headers=viewer_h).json()
    assert s["count"] == 4 and s["total"] == 100 and s["average"] == 25 and s["maximum"] == 40 and s["minimum"] == 10

    s = client.get(
        "/api/v1/analytics/summary",
        params={"category": "stats", "start": "2026-03-02T00:00:00"},
        headers=viewer_h,
    ).json()
    assert s["count"] == 1

    cats = client.get("/api/v1/analytics/by-category", headers=viewer_h).json()
    assert any(c["category"] == "stats" and c["count"] == 4 for c in cats)

    trend = client.get(
        "/api/v1/analytics/trend", params={"category": "stats", "interval": "day"}, headers=viewer_h
    ).json()
    assert [(t["bucket"], t["count"]) for t in trend] == [("2026-03-01", 3), ("2026-03-02", 1)]

    r = client.get("/api/v1/analytics/export", params={"category": "stats"}, headers=viewer_h)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "By Category", "Trend", "Records"]
    assert wb["Records"].max_row == 5


def test_admin_only_endpoints(client, admin_h, user_h):
    for path in ("/api/v1/users", "/api/v1/admin/logs", "/api/v1/admin/db-status", "/api/v1/realtime/history"):
        assert client.get(path, headers=user_h).status_code == 403, path
        assert client.get(path, headers=admin_h).status_code == 200, path


def test_user_role_management(client, admin_h, login):
    client.post("/api/v1/auth/register", json={"username": "carol", "email": "carol@example.com", "password": "secret123"})
    users = client.get("/api/v1/users", params={"search": "carol"}, headers=admin_h).json()
    carol = users["items"][0]
    r = client.patch(f"/api/v1/users/{carol['id']}", json={"role": "viewer"}, headers=admin_h)
    assert r.json()["role"] == "viewer"

    carol_h = login(client, "carol", "secret123")
    assert client.post("/api/v1/records", json={"title": "x", "value": 1, "category": "c"}, headers=carol_h).status_code == 403

    client.patch(f"/api/v1/users/{carol['id']}", json={"is_active": False}, headers=admin_h)
    assert client.get("/api/v1/auth/me", headers=carol_h).status_code == 401

    me = client.get("/api/v1/auth/me", headers=admin_h).json()
    assert client.patch(f"/api/v1/users/{me['id']}", json={"role": "user"}, headers=admin_h).status_code == 400


def test_db_status_and_logs(client, admin_h):
    s = client.get("/api/v1/admin/db-status", headers=admin_h).json()
    assert s["connected"] is True and "users" in s["table_counts"]
    logs = client.get("/api/v1/admin/logs", params={"category": "auth"}, headers=admin_h).json()
    assert logs["total"] >= 1


def test_thresholds(client, admin_h, user_h):
    assert client.put("/api/v1/thresholds/stats", json={"threshold": 25}, headers=user_h).status_code == 403
    r = client.put("/api/v1/thresholds/stats", json={"threshold": 25}, headers=admin_h)
    assert r.status_code == 200
    rec = client.post("/api/v1/records", json={"title": "t", "value": 26, "category": "stats"}, headers=user_h).json()
    assert rec["is_anomaly"] is True
    assert client.delete("/api/v1/thresholds/stats", headers=admin_h).status_code == 204
