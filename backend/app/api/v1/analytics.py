import io
from datetime import datetime
from enum import Enum

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import Integer, case, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import require_perm
from app.core.permissions import Perm
from app.core.timeutil import now, to_naive_local
from app.db.session import engine, get_db
from app.models.realtime import RealtimeData
from app.models.record import DataRecord
from app.models.user import User
from app.schemas.analytics import CategoryStats, SummaryStats, TrendPoint
from app.services.record_filters import RecordFilters, record_filters

router = APIRouter(prefix="/analytics", tags=["Analytics"])

EXPORT_LIMIT = 50_000


class Interval(str, Enum):
    minute = "minute"
    hour = "hour"
    day = "day"
    month = "month"


_FORMATS = {
    Interval.minute: "%Y-%m-%d %H:%i",
    Interval.hour: "%Y-%m-%d %H:00",
    Interval.day: "%Y-%m-%d",
    Interval.month: "%Y-%m",
}


def _bucket(column, interval: Interval):
    """時間分桶運算式（ORM func，依方言選擇對應函式）。"""
    fmt = _FORMATS[interval]
    if engine.dialect.name == "sqlite":
        return func.strftime(fmt.replace("%i", "%M"), column)
    return func.date_format(column, fmt)


def _agg(value_col, anomaly_col):
    return (
        func.count().label("count"),
        func.sum(value_col).label("total"),
        func.avg(value_col).label("average"),
        func.max(value_col).label("maximum"),
        func.min(value_col).label("minimum"),
        func.coalesce(func.sum(case((anomaly_col.is_(True), 1), else_=0)), 0).cast(Integer).label("anomalies"),
    )


def _r(v):
    return round(float(v), 4) if v is not None else None


async def _summary(db: AsyncSession, f: RecordFilters) -> SummaryStats:
    stmt = f.apply(select(*_agg(DataRecord.value, DataRecord.is_anomaly)))
    row = (await db.execute(stmt)).one()
    return SummaryStats(
        count=row.count,
        total=_r(row.total),
        average=_r(row.average),
        maximum=_r(row.maximum),
        minimum=_r(row.minimum),
        anomalies=row.anomalies or 0,
        start=f.start,
        end=f.end,
    )


async def _by_category(db: AsyncSession, f: RecordFilters) -> list[CategoryStats]:
    stmt = f.apply(select(DataRecord.category, *_agg(DataRecord.value, DataRecord.is_anomaly)))
    stmt = stmt.group_by(DataRecord.category).order_by(DataRecord.category)
    return [
        CategoryStats(
            category=r.category,
            count=r.count,
            total=_r(r.total),
            average=_r(r.average),
            maximum=_r(r.maximum),
            minimum=_r(r.minimum),
            anomalies=r.anomalies or 0,
        )
        for r in (await db.execute(stmt)).all()
    ]


async def _trend(db: AsyncSession, f: RecordFilters, interval: Interval, by_category: bool) -> list[TrendPoint]:
    bucket = _bucket(DataRecord.timestamp, interval).label("bucket")
    cols = [bucket]
    if by_category:
        cols.append(DataRecord.category)
    stmt = f.apply(
        select(
            *cols,
            func.count().label("count"),
            func.avg(DataRecord.value).label("average"),
            func.max(DataRecord.value).label("maximum"),
            func.min(DataRecord.value).label("minimum"),
        )
    )
    group = [bucket, DataRecord.category] if by_category else [bucket]
    stmt = stmt.group_by(*group).order_by(*group)
    return [
        TrendPoint(
            bucket=r.bucket,
            category=r.category if by_category else None,
            count=r.count,
            average=_r(r.average),
            maximum=_r(r.maximum),
            minimum=_r(r.minimum),
        )
        for r in (await db.execute(stmt)).all()
    ]


@router.get("/summary", response_model=SummaryStats, summary="統計分析：總計 / 平均 / 最大 / 最小")
async def summary(
    filters: RecordFilters = Depends(record_filters),
    _: User = Depends(require_perm(Perm.ANALYTICS_READ)),
    db: AsyncSession = Depends(get_db),
):
    return await _summary(db, filters)


@router.get("/by-category", response_model=list[CategoryStats], summary="分類資料聚合")
async def by_category(
    filters: RecordFilters = Depends(record_filters),
    _: User = Depends(require_perm(Perm.ANALYTICS_READ)),
    db: AsyncSession = Depends(get_db),
):
    return await _by_category(db, filters)


@router.get("/trend", response_model=list[TrendPoint], summary="資料趨勢（依時間分桶）")
async def trend(
    interval: Interval = Query(Interval.day),
    by_category: bool = Query(True),
    filters: RecordFilters = Depends(record_filters),
    _: User = Depends(require_perm(Perm.ANALYTICS_READ)),
    db: AsyncSession = Depends(get_db),
):
    return await _trend(db, filters, interval, by_category)


@router.get("/realtime/summary", summary="即時資料：各指標統計")
async def realtime_summary(
    start: datetime | None = None,
    end: datetime | None = None,
    _: User = Depends(require_perm(Perm.ANALYTICS_READ)),
    db: AsyncSession = Depends(get_db),
):
    stmt = select(RealtimeData.metric, *_agg(RealtimeData.value, RealtimeData.is_anomaly))
    if start:
        stmt = stmt.where(RealtimeData.timestamp >= to_naive_local(start))
    if end:
        stmt = stmt.where(RealtimeData.timestamp <= to_naive_local(end))
    stmt = stmt.group_by(RealtimeData.metric).order_by(RealtimeData.metric)
    return [
        {
            "metric": r.metric,
            "count": r.count,
            "total": _r(r.total),
            "average": _r(r.average),
            "maximum": _r(r.maximum),
            "minimum": _r(r.minimum),
            "anomalies": r.anomalies or 0,
        }
        for r in (await db.execute(stmt)).all()
    ]


@router.get("/realtime/trend", summary="即時資料：依時間分桶的趨勢")
async def realtime_trend(
    interval: Interval = Query(Interval.minute),
    metric: list[str] | None = Query(None),
    start: datetime | None = None,
    end: datetime | None = None,
    _: User = Depends(require_perm(Perm.ANALYTICS_READ)),
    db: AsyncSession = Depends(get_db),
):
    bucket = _bucket(RealtimeData.timestamp, interval).label("bucket")
    stmt = select(
        bucket,
        RealtimeData.metric,
        func.avg(RealtimeData.value).label("average"),
        func.max(RealtimeData.value).label("maximum"),
        func.min(RealtimeData.value).label("minimum"),
        func.count().label("count"),
    )
    if metric:
        stmt = stmt.where(RealtimeData.metric.in_(metric))
    if start:
        stmt = stmt.where(RealtimeData.timestamp >= to_naive_local(start))
    if end:
        stmt = stmt.where(RealtimeData.timestamp <= to_naive_local(end))
    stmt = stmt.group_by(bucket, RealtimeData.metric).order_by(bucket, RealtimeData.metric)
    return [
        {
            "bucket": r.bucket,
            "metric": r.metric,
            "average": _r(r.average),
            "maximum": _r(r.maximum),
            "minimum": _r(r.minimum),
            "count": r.count,
        }
        for r in (await db.execute(stmt)).all()
    ]


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row in rows:
        ws.append(row)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"


@router.get(
    "/export",
    summary="下載 Excel（明細 + 統計 + 分類 + 趨勢）",
    response_class=StreamingResponse,
    responses={200: {"content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}}}},
)
async def export_excel(
    interval: Interval = Query(Interval.day),
    filters: RecordFilters = Depends(record_filters),
    _: User = Depends(require_perm(Perm.ANALYTICS_EXPORT)),
    db: AsyncSession = Depends(get_db),
):
    records = (
        await db.scalars(filters.apply(select(DataRecord)).order_by(DataRecord.timestamp.desc()).limit(EXPORT_LIMIT))
    ).unique().all()
    summary_stats = await _summary(db, filters)
    cats = await _by_category(db, filters)
    trend_points = await _trend(db, filters, interval, by_category=True)

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(
        wb,
        "Summary",
        ["指標", "值"],
        [
            ["筆數", summary_stats.count],
            ["總計", summary_stats.total],
            ["平均", summary_stats.average],
            ["最大", summary_stats.maximum],
            ["最小", summary_stats.minimum],
            ["異常筆數", summary_stats.anomalies],
            ["起始時間", str(filters.start or "")],
            ["結束時間", str(filters.end or "")],
            ["分類", ", ".join(filters.categories or [])],
            ["匯出時間", now().isoformat(sep=" ")],
        ],
    )
    _write_sheet(
        wb,
        "By Category",
        ["分類", "筆數", "總計", "平均", "最大", "最小", "異常筆數"],
        [[c.category, c.count, c.total, c.average, c.maximum, c.minimum, c.anomalies] for c in cats],
    )
    _write_sheet(
        wb,
        "Trend",
        ["時間", "分類", "筆數", "平均", "最大", "最小"],
        [[t.bucket, t.category, t.count, t.average, t.maximum, t.minimum] for t in trend_points],
    )
    _write_sheet(
        wb,
        "Records",
        ["ID", "標題", "數值", "分類", "時間", "異常", "建立者", "描述"],
        [
            [r.id, r.title, r.value, r.category, r.timestamp, "Y" if r.is_anomaly else "", r.owner_username, r.description]
            for r in records
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"analytics_{now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
